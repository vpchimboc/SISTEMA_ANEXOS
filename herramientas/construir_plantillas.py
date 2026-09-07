#!/usr/bin/env python3
"""
Convierte los anexos ORIGINALES (documentos ya llenados del proyecto anterior)
en PLANTILLAS Jinja2 para docxtpl, conservando intacto el formato de Word.

Por qué así y no editando a mano en Word: el proceso queda reproducible y
auditable. Si mañana la Coordinación cambia el formato de un anexo, se vuelve
a correr este script sobre el anexo nuevo y las reglas dicen exactamente qué
se sustituyó. Cada regla que no encuentre su texto se reporta como AVISO: es
la señal de que la plantilla cambió y hay que revisarla.

Uso:
    python herramientas/construir_plantillas.py \
        --origen "ruta/a/PLANTILLAS/ANEXOS" \
        --destino plantillas
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

import docx

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from core import docx_utils as du  # noqa: E402
from core import logos as mod_logos  # noqa: E402

# --------------------------------------------------------------------------
# Catálogo de logos: se identifica cada imagen del encabezado por su tamaño en
# bytes, porque la posición izquierda/derecha NO es la misma en todos los
# anexos (en unos el TEC.AZUAY va primero, en otros el de la SENESCYT).
# --------------------------------------------------------------------------

HUELLA_LOGOS = {
    49557: "institucional",   # TEC.AZUAY - Instituto Universitario
    9651: "institucional",    # misma marca, versión reducida
    9052: "senescyt",
    9060: "senescyt",
    9051: "senescyt",
    4440: "senescyt",         # versión horizontal, solo texto
    86653: "institucional",   # TEC.AZUAY en alta resolución (informes ISTA)
    81218: "institucional",
}


def identificar_logos(ruta_docx: Path) -> list[dict]:
    """
    Devuelve los slots de logo con su identidad lógica.

    Normalmente están en el encabezado. Si el documento no tiene encabezado —el
    Anexo 5 es así— se buscan en el cuerpo, pero solo se aceptan las imágenes
    cuya huella coincide con un logo conocido: de lo contrario una fotografía
    de evidencia acabaría tratada como logo y se reemplazaría.
    """
    imgs = mod_logos.inspeccionar(ruta_docx)
    en_cuerpo = False
    if not imgs:
        imgs = [i for i in mod_logos.inspeccionar(ruta_docx, incluir_cuerpo=True)
                if i["bytes_tam"] in HUELLA_LOGOS]
        en_cuerpo = bool(imgs)

    slots = []
    for i, img in enumerate(imgs):
        slots.append({
            "indice": i,
            "logico": HUELLA_LOGOS.get(img["bytes_tam"], "desconocido"),
            "ancho_cm": img["ancho_cm"],
            "alto_cm": img["alto_cm"],
            "media": img["media"],
            "en_cuerpo": en_cuerpo,
        })
    return slots


# --------------------------------------------------------------------------
# Reglas de sustitución por anexo
#
# El orden importa: primero las cadenas más largas y específicas, porque una
# regla corta aplicada antes puede destruir el contexto que necesita otra.
# --------------------------------------------------------------------------

NOMBRE_PROYECTO_V1 = (
    "Mantenimiento de la Infraestructura Tecnológica y Red de datos de "
    "Laboratorios de Computación de las Unidades Educativas Fiscales de la "
    "Coordinación de Educación Zonal 6” - Unidad Educativa Luis Monsalve Pozo - Fase V"
)
NOMBRE_PROYECTO_V2 = (
    "Mantenimiento de la infraestructura tecnológica y red de datos de "
    "laboratorios de computación   de las unidades educativas de la "
    "Coordinación de Educación Zonal 6” - Fase V Unidad Educativa Luis Monsalve Pozo"
)
CARRERA_LARGA_1 = "Tecnología Superior Administración de Infraestructura y Plataformas Tecnológicas"
CARRERA_LARGA_2 = "Tecnología Superior en Administración de Infraestructura y Plataformas Tecnológicas"

# Reglas comunes a casi todos los anexos. Son OPCIONALES: cada documento usa
# una sola de las dos redacciones del nombre del proyecto, así que es normal
# que la otra no aparezca y no debe reportarse como problema.
# Las comillas del original están descuadradas (hay una ” en medio del nombre
# del proyecto y a veces falta la de apertura). Se absorben las comillas que
# haya alrededor y se emite siempre el par completo.
COMUNES = [
    ("re:“?" + re.escape(NOMBRE_PROYECTO_V1) + "”?", "“{{ p.nombre }}”"),
    ("re:“?" + re.escape(NOMBRE_PROYECTO_V2) + "”?", "“{{ p.nombre }}”"),
    (CARRERA_LARGA_2, "{{ p.carrera }}"),
    (CARRERA_LARGA_1, "{{ p.carrera }}"),
]

OPCIONALES = {b for b, _ in COMUNES}

REGLAS: dict[str, list[tuple[str, str]]] = {}

# ---- ANEXO 1: Delegación de docentes (uno por docente) --------------------
REGLAS["ANEXO 1"] = COMUNES + [
    ("Cuenca, 22  de septiembre de 2025", "{{ p.ciudad }}, {{ f.anexo1.larga }}"),
    ("Magister,", "{{ docente.tratamiento }},"),
    ("DELEGADO como DOCENTE DE APOYO en el proyecto",
     "DELEGADO como {{ docente.cargo }} en el proyecto"),
    ("Ing. Mónica Fernanda Galarza Rodas Mgtr.", "{{ coord_carrera.firma }}"),
    ("COORDINADORA DE LA CARRERA TAIPT", "{{ coord_carrera.cargo }}"),
    ("Mónica Galarza Rodas", "{{ docente.nombre }}"),
    ("DOCENTE DE APOYO", "{{ docente.cargo }}"),
    ("Fecha: 22/09/2025", "Fecha: {{ f.anexo1.corta }}"),
]

# ---- ANEXO 2: Convocatoria (uno por proyecto) ----------------------------
REGLAS["ANEXO 2"] = COMUNES + [
    ("CONVOCATORIA – TAIPT – 2025 – 02", "{{ p.codigo_convocatoria }}"),
    ("Cuenca, 23 de septiembre de 2025", "{{ p.ciudad }}, {{ f.anexo2.larga }}"),
    ("estudiantes de cuarto ciclo en adelante de la",
     "estudiantes de {{ p.ciclo_minimo }} ciclo en adelante de la"),
    ("que se desarrollará en conjunto con la Unidad Educativa Luis Monsalve Pozo,",
     "que se desarrollará en conjunto con la {{ p.entidad }},"),
    ("La fecha máxima en la que se receptarán las solicitudes es el 14 de agosto de 2025.",
     "La fecha máxima en la que se receptarán las solicitudes es el {{ f.anexo2_limite.larga }}."),
    ("Ing. Jonnathan Fernando Nivicela Arbito (jonnathan.nivicela@ucuenca.edu.ec)",
     "{{ responsable.firma }} ({{ responsable.email }})"),
    ("Ing. Jonnathan Fernando Nivicela Arbito,", "{{ responsable.firma }},"),
    ("CARRERA TAIPT", "CARRERA {{ p.carrera_corta }}"),
]

# ---- ANEXO 3: Solicitud del estudiante (uno por estudiante) ---------------
REGLAS["ANEXO 3"] = COMUNES + [
    ("Cuenca, 24 de septiembre de 2025", "{{ p.ciudad }}, {{ f.anexo3.larga }}"),
    ("Ingeniero", "{{ responsable.tratamiento }}"),
    ("Jonnathan Fernando Nivicela Arbito", "{{ responsable.nombre }}"),
    ("semestre 30-2025-2P", "semestre {{ p.periodo_academico }}"),
    ("paralelo “A”", "paralelo “{{ estudiante.paralelo }}”"),
    ("jornada Matutina", "jornada {{ estudiante.jornada }}"),
    ("Molina Ortiz Jostin Alexander", "{{ estudiante.nombre_completo }}"),
    ("0151181898", "{{ estudiante.cedula }}"),
]

# ---- ANEXO 4: Respuesta al estudiante (uno por estudiante) ----------------
REGLAS["ANEXO 4"] = COMUNES + [
    ("Cuenca, 25 de septiembre de 2025", "{{ p.ciudad }}, {{ f.anexo4.larga }}"),
    ("Señor", "{{ estudiante.tratamiento }}"),
    ("Estudiante de la carrera TAIPT", "Estudiante de la carrera {{ p.carrera_corta }}"),
    ("ha designado como director del proyecto al Ing. Williams Trelles Ávila",
     "ha designado como director del proyecto al {{ director.firma }}"),
    ("del proyecto es Mgtr. Carlos Gonzalo Morales Figueroa.",
     "del proyecto es {{ representante.firma }}."),
    ("Al finalizar las 100 horas", "Al finalizar las {{ p.total_horas }} horas"),
    ("Ing. Jonnathan Fernando Nivicela Arbito", "{{ responsable.firma }}"),
    ("Carrera de TAIPT", "Carrera de {{ p.carrera_corta }}"),
    ("Fecha: 25 de septiembre de 2025", "Fecha: {{ f.anexo4.larga }}"),
    ("Jostin Alexander Molina Ortiz", "{{ estudiante.nombre_completo }}"),
]

# ---- ANEXO 5: Delegación de estudiantes a docente de apoyo ---------------
REGLAS["ANEXO 5"] = COMUNES + [
    ("Cuenca, 25 de septiembre de 2025", "{{ p.ciudad }}, {{ f.anexo5.larga }}"),
    ("Magíster", "{{ docente.tratamiento }}"),
    ("Mónica Galarza Rodas", "{{ docente.nombre }}"),
    ("Ing. Jonnathan Fernando Nivicela Arbito", "{{ responsable.firma }}"),
    ("RESPONSABLE DE PRÁCTICAS PRE PROFESIONALES DE SERVICIO COMUNITARIO TAIPT",
     "RESPONSABLE DE PRÁCTICAS PRE PROFESIONALES DE SERVICIO COMUNITARIO {{ p.carrera_corta }}"),
    ("Fecha: 25 de septiembre de 2025", "Fecha: {{ f.anexo5.larga }}"),
]


# --------------------------------------------------------------------------
# Transformaciones estructurales (tablas que se vuelven bucles)
# --------------------------------------------------------------------------

def _entre(doc, desde: str, hasta: str) -> list:
    """Párrafos con texto que están entre dos párrafos ancla (sin incluirlos)."""
    seleccion, dentro = [], False
    for p in doc.paragraphs:
        texto = p.text.strip()
        if not dentro:
            if desde in texto:
                dentro = True
            continue
        if hasta in texto:
            break
        if texto:
            seleccion.append(p)
    return seleccion


def estructura_anexo2(doc) -> None:
    """Las tres listas de la convocatoria pasan a ser bucles."""
    # 1. Actividades a desarrollar (viene del plan de aprendizaje del proyecto).
    du.lista_a_bucle(
        _entre(doc, "Las actividades a desarrollar son", "deberán haber aprobado"),
        "for a in actividades", "{{ a.descripcion }}",
    )
    # 2. Asignaturas requisito.
    du.lista_a_bucle(
        _entre(doc, "deberán haber aprobado las siguientes asignaturas",
               "La fecha máxima"),
        "for r in requisitos", "{{ r.nombre }}",
    )
    # 3. Cronograma del proceso de selección.
    du.tabla_a_bucle(
        doc.tables[0],
        fila_modelo=1,
        expresion="for c in cronograma",
        celdas={0: "{{ c.actividad }}", 1: "{{ c.fecha }}"},
        filas_a_borrar=[2, 3, 4],
    )


def estructura_anexo5(doc) -> None:
    """Tabla de estudiantes delegados -> bucle sobre `estudiantes`."""
    tabla = doc.tables[0]
    du.tabla_a_bucle(
        tabla,
        fila_modelo=1,
        expresion="for e in estudiantes",
        celdas={0: "{{ e.nombre_completo }}", 1: "{{ e.cedula }}"},
        filas_a_borrar=[2, 3],
    )


# ---- ANEXO 6: Plan de aprendizaje (uno por estudiante) --------------------
REGLAS["ANEXO 6"] = COMUNES + [
    ("Ing. Mónica Galarza Rodas", "{{ estudiante.docente_apoyo.firma }}"),
    ("Unidad Educativa Luis Monsalve Pozo", "{{ p.entidad }}"),
    ("José Alfredo Murillo Quizhpi", "{{ estudiante.nombre_completo }}"),
    ("30-2025-2P", "{{ p.periodo_academico }}"),
    ("26 de septiembre de 2025", "{{ f.anexo6.larga }}"),
    ("Ing. Jonnathan Fernando Nivicela Arbito", "{{ responsable.firma }}"),
    ("Responsable de Vinculación de la carrera TAIPT",
     "Responsable de Vinculación de la carrera {{ p.carrera_corta }}"),
    ("Mgtr. Marilyn Salazar", "{{ coord_vinculacion.firma }}"),
]

# ---- ANEXOS 6.1 y 6.2: seguimiento del plan (uno por estudiante) ----------
_REGLAS_6X = COMUNES + [
    ("Ing. Mónica Galarza Rodas", "{{ estudiante.docente_apoyo.firma }}"),
    ("Ing. Williams Trelles Ávila", "{{ director.firma }}"),
]
REGLAS["ANEXO 6.1"] = _REGLAS_6X + [
    ("Estudiante: Alexander Miguel Coloma Ortiz", "Estudiante: {{ estudiante.nombre_completo }}"),
    ("Fecha: 20/10/2025", "Fecha: {{ f.anexo6_1.corta }}"),
]
REGLAS["ANEXO 6.2"] = _REGLAS_6X + [
    ("Estudiante: José Alfredo Murillo Quizhpi", "Estudiante: {{ estudiante.nombre_completo }}"),
    ("Fecha: 15/12/2025", "Fecha: {{ f.anexo6_2.corta }}"),
]

# ---- ANEXO 7: Planificación mensual (uno por mes) ------------------------
REGLAS["ANEXO 7"] = COMUNES + [
    ("Empresa/Institución: Unidad Educativa Luis Monsalve Pozo",
     "Empresa/Institución: {{ p.entidad }}"),
    ("Director del proyecto: Mgtr. Williams Hidalgo Trelles Avila",
     "Director del proyecto: {{ director.firma }}"),
    ("Fecha Planificación: 26/09/2025", "Fecha Planificación: {{ mes.fecha_planificacion_corta }}"),
    ("Mes/año Planificado: Septiembre 2025", "Mes/año Planificado: {{ mes.etiqueta }}"),
    ("Mgtr. Williams Hidalgo Trelles Avila", "{{ director.firma }}"),
]

# ---- ANEXO 8: Registro diario (uno por estudiante) -----------------------
REGLAS["ANEXO 8"] = COMUNES + [
    ("Entidad Beneficiaria: Unidad Educativa Luis Monsalve Pozo, Coordinación de Educación Zonal 6",
     "Entidad Beneficiaria: {{ p.entidad }}"),
    ("re:Nombre de Estudiante:\\s*Jose Alfredo Murillo Quizhpi",
     "Nombre de Estudiante: {{ estudiante.nombre_completo }}"),
    ("Nro de Cédula: 0107006686", "Nro de Cédula: {{ estudiante.cedula }}"),
    ("Murillo Quizhpi Jose Alfredo", "{{ estudiante.nombre_apellido }}"),
    ("Mgtr. Carlos Gonzalo Morales Figueroa", "{{ representante.firma }}"),
    ("Mgtr. Mónica Galarza Rodas", "{{ estudiante.docente_apoyo.firma }}"),
    ("Mgtr. Williams Hidalgo Trelles Ávila", "{{ director.firma }}"),
]

# ---- ANEXO 9: Seguimiento mensual (uno por mes y docente de apoyo) -------
REGLAS["ANEXO 9"] = COMUNES + [
    ("Empresa/Institución: Unidad Educativa Luis Monsalve Pozo",
     "Empresa/Institución: {{ p.entidad }}"),
    ("re:Mes:\\s*Septiembre 2025", "Mes: {{ mes.etiqueta }}"),
    ("Fecha de Seguimiento: 01/10/2025", "Fecha de Seguimiento: {{ mes.fecha_seguimiento_corta }}"),
    ("Williams Trelles Ávila", "{{ docente.nombre }}"),
    ("Mgtr. Williams Hidalgo Trelles Avila", "{{ director.firma }}"),
]

# ---- ANEXO 11: Evaluación al estudiante (uno por estudiante) -------------
REGLAS["ANEXO 11"] = COMUNES + [
    ("re:Nombres Completos:\\s*Jose Alfredo Murillo Quizhpi",
     "Nombres Completos:  {{ estudiante.nombre_completo }}"),
    ("Nº de Cédula: 0107006686", "Nº de Cédula: {{ estudiante.cedula }}"),
    ("E- Mail: josea.murillo.est@tecazuay.edu.ec", "E- Mail: {{ estudiante.email }}"),
    ("Carrera: Administración de Infraestructura y Plataformas Tecnológicas",
     "Carrera: {{ p.carrera }}"),
    ("Fecha de Inicio: 29/09/2025", "Fecha de Inicio: {{ p.fecha_inicio_corta }}"),
    ("Fecha de Finalización: 12/12/2025", "Fecha de Finalización: {{ p.fecha_fin_corta }}"),
    ("re:Número Total de Horas:(\\s*)100", "Número Total de Horas:\\g<1>{{ p.total_horas }}"),
    ("Fecha de Evaluación: 15/12/2025", "Fecha de Evaluación: {{ p.fecha_evaluacion_corta }}"),
    ("Obteniendo un promedio de 20/20, con lo que APRUEBA",
     "Obteniendo un promedio de {{ ev.promedio }}, con lo que {{ ev.resultado }}"),
    ("Mgtr. Mónica Galarza Rodas", "{{ estudiante.docente_apoyo.firma }}"),
    ("Mgtr. Williams Hidalgo Trelles Ávila", "{{ director.firma }}"),
]

# ---- ANEXO 12: Registro de beneficiarios (uno por jornada) ---------------
REGLAS["ANEXO 12"] = COMUNES + [
    ("re:FECHA:\\s*05/12/2025", "FECHA:  {{ jornada.fecha }}"),
    ("3 horas", "{{ jornada.horas }}"),
    ("ENTIDAD BENEFICIARIA:  Unidad Educativa Luis Monsalve Pozo, Coordinación de Educación Zonal 6",
     "ENTIDAD BENEFICIARIA:  {{ p.entidad }}"),
    ("REPRESENTANTE: Mgtr. Carlos Gonzalo Morales Figueroa",
     "REPRESENTANTE: {{ representante.firma }}"),
    ("TELÉFONO: 0985991161", "TELÉFONO: {{ p.entidad_telefono }}"),
    ("E MAIL: ue.luismonsalve@gmail.com", "E MAIL: {{ p.entidad_email }}"),
    ("ASUNTO: Primera Jornada de Capacitación", "ASUNTO: {{ jornada.asunto }}"),
    ("Mgtr. Carlos Gonzalo Morales Figueroa", "{{ representante.firma }}"),
    ("Mgtr. Williams Hidalgo Trelles Ávila", "{{ director.firma }}"),
]

# ---- ANEXO 13: Registro de visitas (uno por docente de apoyo) ------------
REGLAS["ANEXO 13"] = COMUNES + [
    ("Mgtr. Williams Trelles Ávila", "{{ docente.firma }}"),
    ("Mgtr. Carlos Gonzalo Morales Figueroa", "{{ representante.firma }}"),
    ("30-2025-2P", "{{ p.periodo_academico }}"),
    ("Unidad Educativa Luis Monsalve Pozo, Coordinación de Educación Zonal 6",
     "{{ p.entidad }}"),
    ("re:^Cuarto$", "{{ p.ciclo }}"),
]

# ---- ANEXO 10: Informe de culminación (uno por estudiante) ---------------
REGLAS["ANEXO 10"] = COMUNES + [
    ("Cuenca, 27/05/2024", "{{ p.ciudad }}, {{ p.fecha_fin_corta }}"),
    ("Mgtr. Williams Hidalgo Trelles Avila", "{{ director.firma }}"),
    ("Calle Virgen  del Rosario y Santa Catalina", "{{ p.entidad_direccion }}"),
    ("Mgtr. Mónica Fernanda Galarza Rodas", "{{ estudiante.docente_apoyo.firma }}"),
    ("Mgtr. Mónica Galarza Rodas", "{{ estudiante.docente_apoyo.firma }}"),
    ("Ismael Alejandro Sevillano Guachichullca", "{{ estudiante.nombre_completo }}"),
    ("Mgtr. Carlos Gonzalo Morales Figueroa", "{{ representante.firma }}"),
    ("100 horas", "{{ p.total_horas }} horas"),
    ("18/08/2025", "{{ p.fecha_inicio_corta }}"),
    ("31/10/2025", "{{ p.fecha_fin_corta }}"),
]


# --------------------------------------------------------------------------
# Estructuras de los anexos con tablas repetitivas
# --------------------------------------------------------------------------

def estructura_anexo6(doc) -> None:
    """Plan de aprendizaje: la tabla de actividades pasa a ser bucle."""
    tabla = doc.tables[1]
    du.tabla_a_bucle(
        tabla, fila_modelo=1, expresion="for a in actividades",
        celdas={0: "{{ a.nro }}", 1: "{{ a.descripcion }}",
                2: "{{ a.asignatura }}", 3: "{{ a.resultados_aprendizaje }}",
                4: "{{ a.horas }}h"},
        filas_a_borrar=list(range(2, 9)),
    )
    # La fila del total queda fuera del bucle y se calcula.
    total = tabla.rows[-1]
    du.reemplazar_texto_celda(total.cells[-1], "{{ p.total_horas }}h")


def _estructura_6x(doc) -> None:
    """Seguimiento del plan: una fila por actividad evaluada."""
    du.tabla_a_bucle(
        doc.tables[0], fila_modelo=2, expresion="for a in plan",
        celdas={0: "{{ a.nro }}", 1: "{{ a.descripcion }}", 2: "{{ a.avance }}",
                3: "{{ a.poco }}", 4: "{{ a.satisfactorio }}", 5: "{{ a.muy }}",
                6: "{{ a.si }}", 7: "{{ a.no }}"},
        filas_a_borrar=list(range(3, 10)),
    )


def estructura_anexo7(doc) -> None:
    """Planificación mensual: horas docentes, horas estudiantes y firmas."""
    du.tabla_a_bucle(
        doc.tables[0], fila_modelo=1, expresion="for r in plan_docentes",
        celdas={0: "{{ r.resultado }}", 1: "{{ r.actividad }}", 2: "{{ r.nombre }}",
                3: "{{ r.cedula }}", 4: "{{ r.horas }}", 5: "{{ r.fecha_inicio }}",
                6: "{{ r.fecha_fin }}", 7: "{{ r.observaciones }}"},
        filas_a_borrar=list(range(2, 7)),
    )
    du.tabla_a_bucle(
        doc.tables[1], fila_modelo=1, expresion="for r in plan_estudiantes",
        celdas={0: "{{ r.resultado }}", 1: "{{ r.actividad }}", 2: "{{ r.nombre }}",
                3: "{{ r.cedula }}", 4: "{{ r.horas }}", 5: "{{ r.fecha_inicio }}",
                6: "{{ r.fecha_fin }}", 7: "{{ r.observaciones }}"},
        filas_a_borrar=list(range(2, 16)),
    )
    du.tabla_a_bucle(
        doc.tables[3], fila_modelo=1, expresion="for e in estudiantes",
        celdas={0: "{{ loop.index }}", 1: "{{ e.cedula }}",
                2: "{{ e.nombre_completo }}", 3: "", 4: "{{ mes.fecha_socializacion_corta }}"},
        filas_a_borrar=list(range(2, 6)),
    )


def estructura_anexo8(doc) -> None:
    """Registro diario: una fila por actividad registrada, y el total."""
    tabla = doc.tables[1]
    ultima = len(tabla.rows) - 1
    du.tabla_a_bucle(
        tabla, fila_modelo=1, expresion="for r in registro",
        celdas={0: "{{ r.fecha }}", 1: "{{ r.descripcion }}",
                2: "{{ r.lugar }}", 3: "{{ r.horas }}h"},
        filas_a_borrar=list(range(2, ultima)),
    )
    du.reemplazar_texto_celda(tabla.rows[-1].cells[-1], "{{ total_registro }}h")


def estructura_anexo9(doc) -> None:
    """Se conserva una sola ficha (el original trae una por docente de apoyo)."""
    du.recortar_desde(doc, 22)
    du.tabla_a_bucle(
        doc.tables[0], fila_modelo=1, expresion="for s in seguimiento",
        celdas={0: "{{ loop.index }}", 1: "{{ s.descripcion }}",
                2: "{{ s.estudiante }}", 3: "{{ s.fecha_planificada }}",
                4: "{{ s.finalizada }}", 5: "{{ s.fecha_fin_prevista }}",
                6: "{{ s.avance }}"},
        filas_a_borrar=[2, 3],
    )


def estructura_anexo11(doc) -> None:
    """
    Evaluación: las 'x' se colocan por celda según el puntaje elegido.
    `marca` es una función que se inyecta en el contexto: devuelve 'x' si ese
    ítem tiene ese puntaje. Es más legible en la plantilla que ocho condicionales.
    """
    items = {0: ["1.1", "1.2", "1.3", "1.4"], 1: ["2.1", "2.2", "2.3", "2.4"]}
    for i_tabla, claves in items.items():
        tabla = doc.tables[i_tabla]
        for fila, clave in enumerate(claves, start=1):
            for col in range(1, 6):
                du.reemplazar_texto_celda(
                    tabla.rows[fila].cells[col + 1],
                    "{{ marca('%s', %d) }}" % (clave, col),
                )
        du.reemplazar_texto_celda(tabla.rows[5].cells[-1],
                                  "{{ ev.bloque%d }}" % (i_tabla + 1))


def estructura_anexo12(doc) -> None:
    """Se conserva una sola ficha de registro (una por jornada de capacitación)."""
    du.recortar_desde(doc, 4)


def estructura_anexo13(doc) -> None:
    """Una sola ficha de visitas; la lista de estudiantes y las visitas se repiten."""
    du.recortar_desde(doc, 22)
    du.lista_a_bucle(
        [p for p in doc.paragraphs
         if p.text.strip() in {"Paul Angel Mejia Morocho",
                               "Kevin Israel Brito Dumaguala",
                               "Jostin Alexander Molina Ortiz"}],
        "for e in estudiantes", "{{ e.nombre_completo }}",
    )
    # Cada visita ocupa tres filas: fecha, cabecera y contenido.
    du.tabla_a_bucle_grupo(
        doc.tables[1], desde=0, hasta=2, expresion="for v in visitas",
        celdas_por_fila={
            0: {1: "{{ v.fecha }}", 3: "{{ v.hora_inicio }} - {{ v.hora_fin }}"},
            2: {0: "{{ v.asunto }}", 1: "{{ v.actividades }}",
                2: "{{ v.observaciones }}"},
        },
    )


# Ficha de datos del Anexo 10: se localiza cada fila por su rótulo, no por su
# índice, para que siga funcionando si añaden o quitan filas en el formato.
FICHA_ANEXO_10 = [
    ("DATOS DEL PROYECTO", "NOMBRE:", -1, "{{ p.nombre }}"),
    ("DATOS GENERALES DE LA EMPRESA", "NOMBRE:", -1, "{{ p.entidad }}"),
    ("DATOS DEL DOCENTE DE APOYO", "NOMBRE COMPLETO:", 3,
     "{{ estudiante.docente_apoyo.cedula }}"),
    ("DATOS DEL DOCENTE DE APOYO", "CORREO ELECTRÓNICO:", -1,
     "{{ estudiante.docente_apoyo.email }}"),
    ("DATOS DEL ESTUDIANTE", "CÉDULA DE IDENTIDAD:", -1, "{{ estudiante.cedula }}"),
    ("DATOS DEL ESTUDIANTE", "ÚLTIMO CICLO APROBADO:", -1, "{{ estudiante.ciclo }}"),
    ("DATOS DEL ESTUDIANTE", "CORREO ELECTRÓNICO:", -1, "{{ estudiante.email }}"),
    ("DATOS DEL ESTUDIANTE", "TELÉFONO:", -1, "{{ estudiante.telefono }}"),
    ("DATOS DEL ADMINISTRADOR", "CÉDULA DE IDENTIDAD:", -1, "{{ representante.cedula }}"),
    ("DATOS DEL ADMINISTRADOR", "CORREO ELECTRÓNICO:", -1, "{{ p.entidad_email }}"),
]


def _rellenar_ficha(tabla, ficha) -> list[str]:
    """
    Recorre la tabla llevando cuenta de la sección en la que va (las filas de
    una sola celda son los encabezados de sección) y escribe cada valor.
    Devuelve las reglas que no encontraron su sitio.
    """
    pendientes = list(ficha)
    seccion = ""
    for fila in tabla.rows:
        celdas, vistas = [], []
        for celda in fila.cells:
            if celda._tc in vistas:
                continue
            vistas.append(celda._tc)
            celdas.append(celda)
        if len(celdas) == 1:
            seccion = celdas[0].text.strip()
            continue
        rotulo = celdas[0].text.strip()
        for regla in list(pendientes):
            sec, rot, idx, valor = regla
            if sec in seccion and rotulo.startswith(rot):
                destino = celdas[idx] if idx < len(celdas) else celdas[-1]
                du.reemplazar_texto_celda(destino, valor)
                pendientes.remove(regla)
                break
    return [f"{sec} / {rot}" for sec, rot, _, _ in pendientes]


def estructura_anexo10(doc) -> None:
    """Informe del estudiante: ficha de datos + tabla de actividades."""
    faltantes = _rellenar_ficha(doc.tables[0], FICHA_ANEXO_10)
    if faltantes:
        raise ValueError("Anexo 10: no se localizaron en la ficha -> "
                         + ", ".join(faltantes))

    # La descripción de la institución venía con el texto de un proyecto aún
    # más antiguo (otra unidad educativa). Pasa a ser un campo del proyecto.
    du.reemplazar_texto_celda(doc.tables[1].rows[0].cells[0],
                              "{{ p.entidad }}\n{{ p.entidad_descripcion }}")

    # Las evidencias semanales las escribe el estudiante en Word; se limpia el
    # ejemplo del proyecto anterior para que no se cuele en la entrega.
    for parrafo in doc.paragraphs:
        texto = parrafo.text.strip()
        if texto.startswith("Actividad:") and len(texto) > len("Actividad:"):
            du.escribir_parrafo(parrafo, "Actividad:")
        elif texto and texto.startswith("Observación de los laboratorios"):
            du.escribir_parrafo(parrafo, "")
    du.tabla_a_bucle(
        doc.tables[2], fila_modelo=1, expresion="for a in actividades",
        celdas={0: "{{ a.descripcion }}", 1: "{{ a.detalle_especifico }}",
                2: "{{ a.producto }}"},
        filas_a_borrar=list(range(2, 9)),
    )


# ==========================================================================
# DOCUMENTOS ADICIONALES (sin número de anexo)
#
# Son los cuatro informes que la Constancia lista como "S/N" más la propia
# Constancia. A diferencia de los anexos, aquí la mayor parte es texto
# redactado: se guarda en la base y se edita desde la app, no se calcula.
#
# Cuidado: estos documentos llevan imágenes en el cuerpo (evidencias,
# cronogramas pegados como imagen, firmas escaneadas). Por eso los bloques de
# redacción se sustituyen párrafo a párrafo con `bloque_a_variable`, que
# respeta los párrafos que contienen dibujos, en lugar de vaciar celdas.
# ==========================================================================

PROYECTO_BD = ("Plataforma web de análisis de los datos, aplicando técnicas de "
               "minería de datos y aprendizaje automático para la toma de decisiones "
               "de la fundación Mensajeros de la Paz. Fase 2.")
PROYECTO_BD_MAY = ("PLATAFORMA WEB DE ANÁLISIS DE LOS DATOS, APLICANDO TÉCNICAS DE "
                   "MINERÍA DE DATOS Y APRENDIZAJE DE MÁQUINA, PARA LA TOMA DE "
                   "DECISIONES DE LA FUNDACIÓN MENSAJEROS DE LA PAZ")

# Las firmas aparecen con espaciado y sufijos inconsistentes entre documentos
# ("Mgtr.  Verónica", "Ing. Freddy … Baculima" sin el "Mgtr." final). Se usan
# expresiones regulares tolerantes al espacio en blanco.
FIRMAS_COMUNES = [
    ("re:Mgtr\\.\\s+Verónica Paulina Chimbo Coronel", "{{ director.firma }}"),
    ("re:Ing\\.\\s+Freddy Vinicio Quito Baculima(?:\\.\\s*Mgtr\\.)?",
     "{{ coord_vinculacion.firma }}"),
    ("Ing. Verónica Chimbo, Mgtr.", "{{ director.firma }}"),
    ("Mgst.Verónica Chimbo Coronel", "{{ director.firma }}"),
    ("Verónica Chimbo Coronel", "{{ director.nombre }}"),
    ("Ing. Jéssica Elizabeth Pinos Pinos, Mgrt.", "{{ responsable.firma }}"),
    ("Ing. Jessica Elizabeth Pinos Pinos. Mgtr.", "{{ responsable.firma }}"),
    ("Mgtr. Victoria Marilyn Salazar Piña", "{{ coord_vinculacion.firma }}"),
]

# ---- Informe del proceso de selección -------------------------------------
REGLAS["SELECCION"] = FIRMAS_COMUNES + [
    ("re:" + re.escape(PROYECTO_BD_MAY) + r"\.?\s*FASE 2\.?", "{{ p.nombre_mayusculas }}"),
    ("re:“?" + re.escape(PROYECTO_BD_MAY) + r"\.?”?\s*FASE 2\.?", "“{{ p.nombre_mayusculas }}”"),
    ("CARRERA DE BIG DATA", "CARRERA DE {{ p.carrera_corta }}"),
    ("Carrera: Big Data", "Carrera: {{ p.carrera_corta }}"),
    ("Fecha: 21/05/2026", "Fecha: {{ f.seleccion.corta }}"),
]

# ---- Informes ISTA (seguimiento y final) ----------------------------------
_REGLAS_ISTA = FIRMAS_COMUNES + [
    # El seguimiento lo escribe con "Fase 2." y el final sin fase y con doble
    # punto. Una sola regla flexible cubre ambos.
    ("re:" + re.escape(PROYECTO_BD[:-9]) + r"\.?\.?\s*(?:Fase 2\.?)?", "{{ p.nombre }}"),
    ("Aplicación de tecnologías para responder a las necesidades reales de los "
     "sectores más vulnerables.", "{{ p.programa_vinculacion }}"),
    ("☐Asesoría Técnica", "{{ o.linea_asesoria }}Asesoría Técnica"),
    ("☒Trabajo Comunitario", "{{ o.linea_comunitario }}Trabajo Comunitario"),
    ("☐Capacitación", "{{ o.linea_capacitacion }}Capacitación"),
    ("re:^Tecnología Superior en Big Data$", "{{ p.carrera }}"),
    ("Fundación Mensajeros de La Paz", "{{ p.entidad }}"),
    ("re:^3 Meses$", "{{ p.plazo_ejecucion }}"),
    ("re:Fecha de inicio de Proyecto:\\s*01/06/2026",
     "Fecha de inicio de Proyecto:  {{ p.fecha_inicio_corta }}"),
    ("Fecha de fin planeado: 04/09/2026", "Fecha de fin planeado: {{ p.fecha_fin_corta }}"),
    ("☐Nacional", "{{ o.alcance_nacional }}Nacional"),
    ("☒Provincial", "{{ o.alcance_provincial }}Provincial"),
    ("☐Cantonal", "{{ o.alcance_cantonal }}Cantonal"),
    ("☐Parroquial", "{{ o.alcance_parroquial }}Parroquial"),
    ("☐Institucional", "{{ o.alcance_institucional }}Institucional"),
    ("☐Internacional", "{{ o.alcance_internacional }}Internacional"),
    ("Fortalecer los procesos administrativos y de salud en la Fundación "
     "Mensajeros de la Paz mediante el desarrollo de una aplicación de big data "
     "en la web que facilite el proceso de toma de decisiones. .",
     "{{ n.objetivo_general }}"),
]

REGLAS["SEGUIMIENTO"] = _REGLAS_ISTA + [
    ("Fecha de seguimiento: 15/07/2026",
     "Fecha de seguimiento: {{ p.fecha_seguimiento_corta }}"),
    ("Fecha de Entrega: 15/07/2026", "Fecha de Entrega: {{ p.fecha_entrega_corta }}"),
]

REGLAS["FINAL"] = _REGLAS_ISTA + [
    ("☒Impacto Social", "{{ o.impacto_social }}Impacto Social"),
    ("☐Impacto Científico", "{{ o.impacto_cientifico }}Impacto Científico"),
    ("☐Impacto Económico", "{{ o.impacto_economico }}Impacto Económico"),
    ("☐Impacto Político", "{{ o.impacto_politico }}Impacto Político"),
    ("☐Otro Impacto", "{{ o.impacto_otro }}Otro Impacto"),
    ("https://drive.google.com/drive/folders/1rtAtPpDK3RJxuushaHk_3Y0j2DmMdV3g",
     "{{ p.enlace_anexos }}"),
    ("Fecha de Entrega: 10/09/2026", "Fecha de Entrega: {{ p.fecha_entrega_corta }}"),
]

# ---- Informe de socialización ---------------------------------------------
REGLAS["SOCIALIZACION"] = FIRMAS_COMUNES + [
    ("re:“?" + re.escape(PROYECTO_BD_MAY) + r"\.?”?\s*FASE 2\.?", "“{{ p.nombre_mayusculas }}”"),
    ("Nombre de la carrera: Tecnología Superior en Big Data",
     "Nombre de la carrera: {{ p.carrera }}"),
    ("Nombre del director de carrera: Ing. Priscila Bernal. Mgtr.",
     "Nombre del director de carrera: {{ p.director_carrera }}"),
    ("Fecha: 07/09/2026", "Fecha: {{ f.socializacion.corta }}"),
    ("Fecha:10/09/2026", "Fecha: {{ f.socializacion_aprobacion.corta }}"),
]

# ---- Constancia de recepción de anexos ------------------------------------
REGLAS["CONSTANCIA"] = [
    ("re:DIRECTORA:\\s*Mgtr\\. Verónica Paulina Chimbo Coronel",
     "DIRECTOR/A: {{ director.firma }}"),
] + FIRMAS_COMUNES + [
    ("re:Nombre del proyecto:\\s*“?" + re.escape(PROYECTO_BD_MAY) + r"\.?”?\s*Fase 2\.?",
     "Nombre del proyecto: “{{ p.nombre_mayusculas }}”"),
    ("re:FECHA DE INICIO:\\s*01/06/2026", "FECHA DE INICIO: {{ p.fecha_inicio_corta }}"),
    ("FECHA DE FIN: 31/08/2026", "FECHA DE FIN: {{ p.fecha_fin_corta }}"),
    ("Fecha: 09/09/2026", "Fecha: {{ f.constancia.corta }}"),
    ("Fecha:  09/09/2026", "Fecha: {{ f.constancia.corta }}"),
]


OPCIONALES |= {b for b, _ in FIRMAS_COMUNES}


def estructura_seleccion(doc) -> None:
    """
    El informe de selección tiene toda la redacción dentro de una tabla de una
    sola columna, alternando fila de rótulo y fila de contenido.
    """
    tabla = doc.tables[0]
    bloques = {3: "{{ n.sel_antecedentes }}",
               5: "{{ n.sel_objetivo }}",
               7: "{{ n.sel_desarrollo }}",
               9: "{{ n.sel_evidencias }}"}
    for fila, etiqueta in bloques.items():
        celda = tabla.rows[fila].cells[0]
        du.bloque_a_variable(list(celda.paragraphs), etiqueta)

    # Dentro del bloque "Desarrollo" hay dos TABLAS ANIDADAS que el reemplazo
    # de párrafos no alcanza: el cronograma de la convocatoria y la tabla de
    # resultados con los nombres de los estudiantes. Sin esto, los estudiantes
    # del proyecto anterior se quedaban en el documento.
    anidadas = tabla.rows[7].cells[0].tables
    if len(anidadas) >= 1:
        crono = anidadas[0]
        du.tabla_a_bucle(
            crono, fila_modelo=1, expresion="for c in cronograma",
            celdas={0: "{{ c.actividad }}", 1: "{{ c.fecha }}"},
            filas_a_borrar=list(range(2, len(crono.rows))))
    if len(anidadas) >= 2:
        resultados = anidadas[1]
        du.tabla_a_bucle(
            resultados, fila_modelo=1, expresion="for e in estudiantes",
            celdas={0: "{{ loop.index }}", 1: "{{ e.nombre_completo }}",
                    2: "SELECCIONADO", 3: ""},
            filas_a_borrar=list(range(2, len(resultados.rows))))


def _estructura_ista_comun(doc) -> None:
    """Partes que comparten el informe de seguimiento y el informe final."""
    # Equipo de trabajo.
    du.tabla_a_bucle(
        doc.tables[0], fila_modelo=1, expresion="for d in docentes_informe",
        celdas={0: "{{ d.nombre }}", 1: "{{ d.cedula }}",
                2: "{{ p.carrera_corta }}", 3: "{{ d.horas }}"},
    )
    estudiantes = doc.tables[1]
    columnas = len(estudiantes.columns)
    celdas = {0: "{{ loop.index }}", 1: "{{ e.nombre_completo }}", 2: "{{ e.cedula }}"}
    if columnas >= 6:                      # el informe final añade "Código de Estudiante"
        celdas.update({3: "{{ e.codigo }}", 4: "{{ p.carrera_corta }}", 5: "{{ e.horas }}"})
    else:
        celdas.update({3: "{{ p.carrera_corta }}", 4: "{{ e.horas }}"})
    du.tabla_a_bucle(
        estudiantes, fila_modelo=1, expresion="for e in estudiantes",
        celdas=celdas,
        filas_a_borrar=list(range(2, len(estudiantes.rows))),
    )

    # Objetivos específicos.
    du.lista_a_bucle(
        du.parrafos_entre(doc, "Objetivos Específicos", "Situación al inicio"),
        "for o in objetivos_especificos", "{{ o.texto }}",
    )
    # Bloques de redacción largos.
    du.bloque_a_variable(
        du.parrafos_entre(doc, "Situación al inicio de la ejecución del proyecto",
                          "Situación actual de los beneficiarios"),
        "{{ n.situacion_inicial }}")
    du.bloque_a_variable(
        du.parrafos_entre(doc, "Situación actual de los beneficiarios",
                          "Descripción de las actividades realizadas"),
        "{{ n.situacion_beneficiarios }}")


def estructura_seguimiento(doc) -> None:
    _estructura_ista_comun(doc)
    actividades = doc.tables[2]
    du.tabla_a_bucle(
        actividades, fila_modelo=1, expresion="for a in actividades_informe",
        celdas={0: "{{ a.actividad }}", 1: "{{ a.cumplimiento }}", 2: "{{ a.fecha }}",
                3: "{{ a.responsables }}", 4: "{{ a.evidencia }}",
                5: "{{ a.observaciones }}"},
        filas_a_borrar=list(range(2, len(actividades.rows))),
    )
    du.lista_a_bucle(du.parrafos_entre(doc, "Conclusiones", "Observaciones"),
                     "for c in conclusiones", "{{ c.texto }}")
    du.bloque_a_variable(du.parrafos_entre(doc, "Observaciones", "Fecha de Entrega"),
                         "{{ n.observaciones_informe }}")


def estructura_final(doc) -> None:
    _estructura_ista_comun(doc)
    actividades = doc.tables[2]
    du.tabla_a_bucle(
        actividades, fila_modelo=1, expresion="for a in actividades_informe",
        celdas={0: "{{ a.actividad }}", 1: "{{ a.cumplimiento }}", 2: "{{ a.fecha }}",
                3: "{{ a.responsables }}", 4: "{{ a.evidencia }}",
                5: "{{ a.observaciones }}"},
        filas_a_borrar=list(range(2, len(actividades.rows))),
    )
    du.bloque_a_variable(
        du.parrafos_entre(doc, "Descripción de Impacto Generado", "Indicadores"),
        "{{ n.impacto_descripcion }}")

    indicadores = doc.tables[3]
    du.tabla_a_bucle(
        indicadores, fila_modelo=1, expresion="for i in indicadores",
        celdas={0: "{{ i.nro }}", 1: "{{ i.descripcion }}", 2: "{{ i.tipo }}"},
        filas_a_borrar=list(range(2, len(indicadores.rows))),
    )
    du.lista_a_bucle(
        du.parrafos_entre(doc, "Resultados de los Indicadores de Impacto",
                          "Matriz de verificación"),
        "for r in resultados_indicadores", "{{ r.texto }}")

    matriz = doc.tables[4]
    du.tabla_a_bucle(
        matriz, fila_modelo=2, expresion="for m in matriz_objetivos",
        celdas={0: "{{ m.objetivo }}", 1: "{{ m.indicador }}",
                2: "{{ m.planificado }}", 3: "{{ m.obtenido }}",
                4: "{{ m.observaciones }}"},
        filas_a_borrar=list(range(3, len(matriz.rows))),
    )
    du.lista_a_bucle(
        du.parrafos_entre(doc, "Resultados alcanzados", "Observaciones"),
        "for x in productos", "{{ x.texto }}")
    du.bloque_a_variable(du.parrafos_entre(doc, "Observaciones", "CONCLUSIONES Y RECOMENDACIONES"),
                         "{{ n.observaciones_informe }}")
    du.lista_a_bucle(du.parrafos_entre(doc, "Conclusiones", "Recomendaciones"),
                     "for c in conclusiones", "{{ c.texto }}")
    du.lista_a_bucle(du.parrafos_entre(doc, "Recomendaciones", "ANEXOS"),
                     "for r in recomendaciones", "{{ r.texto }}")


def estructura_socializacion(doc) -> None:
    """Ficha de divulgación: cuatro secciones libres dentro de una tabla."""
    tabla = doc.tables[0]
    bloques = {3: "{{ n.soc_antecedentes }}",
               5: "{{ n.soc_objetivo }}",
               7: "{{ n.soc_desarrollo }}",
               9: "{{ n.soc_conclusiones }}"}
    for fila, etiqueta in bloques.items():
        if fila < len(tabla.rows):
            du.bloque_a_variable(list(tabla.rows[fila].cells[0].paragraphs), etiqueta)


def estructura_constancia(doc) -> None:
    """Se conserva una sola ficha; el checklist pasa a ser un bucle."""
    du.recortar_desde(doc, 7)
    tabla = doc.tables[0]
    du.tabla_a_bucle(
        tabla, fila_modelo=1, expresion="for c in constancia",
        celdas={0: "{{ c.numero }}", 1: "{{ c.nombre }}",
                2: "{{ c.presentacion }}", 3: "{{ c.marca }}"},
        filas_a_borrar=list(range(2, len(tabla.rows))),
    )


ESTRUCTURA = {
    "SELECCION": estructura_seleccion,
    "SEGUIMIENTO": estructura_seguimiento,
    "FINAL": estructura_final,
    "SOCIALIZACION": estructura_socializacion,
    "CONSTANCIA": estructura_constancia,
    "ANEXO 2": estructura_anexo2,
    "ANEXO 5": estructura_anexo5,
    "ANEXO 6": estructura_anexo6,
    "ANEXO 6.1": _estructura_6x,
    "ANEXO 6.2": _estructura_6x,
    "ANEXO 7": estructura_anexo7,
    "ANEXO 8": estructura_anexo8,
    "ANEXO 9": estructura_anexo9,
    "ANEXO 10": estructura_anexo10,
    "ANEXO 11": estructura_anexo11,
    "ANEXO 12": estructura_anexo12,
    "ANEXO 13": estructura_anexo13,
}


# --------------------------------------------------------------------------
# Archivos de origen por anexo
# --------------------------------------------------------------------------

ORIGENES = {
    "ANEXO 1": "ANEXO 1/ANEXO 1.docx",
    "ANEXO 2": "ANEXO 2/ANEXO 2_Convocatoria.docx",
    "ANEXO 3": "ANEXO 3/ANEXO 3.docx",
    "ANEXO 4": "ANEXO 4/ANEXO 4.docx",
    "ANEXO 5": "ANEXO 5/ANEXO 5.docx",
    "ANEXO 6": "ANEXO 6/ANEXO 6.docx",
    "ANEXO 6.1": "ANEXO 6/ANEXO 6.1.docx",
    "ANEXO 6.2": "ANEXO 6/ANEXO 6.2.docx",
    "ANEXO 7": "ANEXO 7/ANEXO 7_PlanificacionMes 01.docx",
    "ANEXO 8": "ANEXO 8/ANEXO 8.docx",
    "ANEXO 9": "ANEXO 9/ANEXO 9_Seguimiento Planificacion Mes 01.docx",
    "ANEXO 10": "ANEXO 10/ANEXO 10_Informe_Plantilla.docx",
    "ANEXO 11": "ANEXO 11/ANEXO 11.docx",
    "ANEXO 12": "ANEXO 12/ANEXO 12_RegistroBeneficiarios_Apoyo1.docx",
    "ANEXO 13": "ANEXO 13/ANEXO 13.docx",
    "SELECCION": "2 Informes Proceso Seleccion (RP).docx",
    "SEGUIMIENTO": "3 ISTA_Formato_proyectos_vinculacion_informe_seguimiento.docm.docx",
    "FINAL": "4 ISTA_Formato_proyectos_vinculacion_informe_final.docm.docx",
    "SOCIALIZACION": "5 Informe de Socialización.docx",
    "CONSTANCIA": "Constancia anexos big data.docx",
}


# --------------------------------------------------------------------------
# Motor
# --------------------------------------------------------------------------

def construir(origen: Path, destino: Path, solo: list[str] | None = None) -> dict:
    destino.mkdir(parents=True, exist_ok=True)
    manifiesto = {}
    problemas = []

    for anexo, relativo in ORIGENES.items():
        if solo and anexo not in solo:
            continue
        ruta_origen = origen / relativo
        if not ruta_origen.exists():
            problemas.append(f"{anexo}: no se encontró {ruta_origen}")
            continue

        salida = destino / f"{anexo.lower().replace(' ', '_')}.docx"
        shutil.copy2(ruta_origen, salida)

        doc = docx.Document(salida)

        # 1. Transformación estructural (tablas -> bucles), antes del texto,
        #    porque borra filas cuyos textos ya no interesa sustituir.
        if anexo in ESTRUCTURA:
            ESTRUCTURA[anexo](doc)

        # 2. Sustitución de textos.
        conteo = du.reemplazar(doc, REGLAS[anexo])
        sin_efecto = [b for b, n in conteo.items() if n == 0 and b not in OPCIONALES]
        if sin_efecto:
            problemas.append(f"{anexo}: reglas sin coincidencia -> {sin_efecto}")

        doc.save(salida)

        manifiesto[anexo] = {
            "archivo": salida.name,
            "origen": relativo,
            "logos": identificar_logos(salida),
            "reglas_aplicadas": {k: v for k, v in conteo.items() if v},
        }
        print(f"  ✓ {anexo:9s} -> {salida.name}  ({sum(conteo.values())} sustituciones)")

    (destino / "manifiesto.json").write_text(
        json.dumps(manifiesto, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return {"manifiesto": manifiesto, "problemas": problemas}


def _buscar_origen() -> Path | None:
    """
    Busca la carpeta de anexos oficiales sin que haya que escribir la ruta.
    Mira junto al sistema y un par de niveles hacia arriba, que es donde suele
    estar si el sistema se instaló dentro de la carpeta del proyecto.
    """
    candidatos = [
        RAIZ / "PLANTILLAS" / "ANEXOS",
        RAIZ.parent / "PLANTILLAS" / "ANEXOS",
        RAIZ.parent.parent / "PLANTILLAS" / "ANEXOS",
    ]
    for candidato in candidatos:
        if (candidato / "ANEXO 1" / "ANEXO 1.docx").exists():
            return candidato
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--origen", help="Carpeta PLANTILLAS/ANEXOS original")
    ap.add_argument("--destino", default=str(RAIZ / "plantillas"))
    ap.add_argument("--solo", nargs="*", help="Construir solo estos anexos")
    args = ap.parse_args()

    origen = Path(args.origen) if args.origen else _buscar_origen()
    if origen is None:
        print("No se encontró la carpeta PLANTILLAS/ANEXOS. Indícala con --origen.")
        return 2

    print(f"Construyendo plantillas desde: {origen}")
    resultado = construir(origen, Path(args.destino), args.solo)

    if resultado["problemas"]:
        print("\nAVISOS:")
        for x in resultado["problemas"]:
            print("  !", x)
        return 1
    print("\nSin avisos. Todas las reglas encontraron su texto.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
